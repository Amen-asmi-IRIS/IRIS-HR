"""The sick-leave comment and the certificate name reach only the requester and HR (review focus 9)."""

from __future__ import annotations

import json
import logging
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse

from apps.audit.models import AuditEvent
from apps.notifications.models import Notification
from apps.requests import services as request_services
from apps.requests.services import RequestData
from tests.hardening.helpers import MARKER, decide, give_balance, make_request, pdf
from tests.matrix import P

MAY_SEE = {P.OWNER, P.HR}
EVERYONE = [P.OWNER, P.PEER, P.MANAGER, P.OTHER_MANAGER, P.OUTSIDER, P.HR, P.EXEC, P.SYSADMIN]
STAGES = ["pending", "approved", "corrected", "reclassified"]
PERIOD = {"from": "2026-11-01", "to": "2026-11-30"}
# What each report of plan 06 asks for (its `ReportQuery`): a period, a date (`on`), a year, or nothing.
REPORT_QUERIES = {
    "absences": PERIOD,
    "balances": {"on": "2026-11-30"},
    "waiting": {},
    "permission-hours": {"year": "2026"},
    "remote-days": {"year": "2026"},
    "certificates": {},
}


def _text(response) -> str:
    """The whole response as searchable text: headers, JSON body, or every cell of a workbook."""
    if getattr(response, "streaming", False):
        body = b"".join(response.streaming_content)
    else:
        body = response.content
    headers = "\n".join(f"{key}: {value}" for key, value in response.headers.items())
    if body[:2] == b"PK":
        workbook = openpyxl.load_workbook(BytesIO(body), read_only=True)
        cells = [
            str(cell)
            for sheet in workbook
            for row in sheet.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        ]
        return headers + "\n" + "\n".join(workbook.sheetnames) + "\n" + "\n".join(cells)
    return headers + "\n" + body.decode("utf-8", errors="replace")


def _detail_urls(world, request, attachment):
    """Places where the requester and HR are entitled to the comment or the file name."""
    rid = {"public_id": request.public_id}
    return [
        ("request", reverse("requests:request-detail", kwargs=rid), {}),
        ("request attachments", reverse("requests:request-attachments", kwargs=rid), {}),
        ("request history", reverse("requests:request-history", kwargs=rid), {}),
        ("request amendments", reverse("requests:request-amendments", kwargs=rid), {}),
        ("download", reverse("files:download", kwargs={"public_id": attachment.public_id}), {}),
        ("requests mine", reverse("requests:request-list"), {"scope": "mine"}),
        ("requests team", reverse("requests:request-list"), {"scope": "team"}),
        ("requests all", reverse("requests:request-list"), {"scope": "all"}),
        ("requests decided", reverse("requests:request-list"), {"scope": "decided"}),
        ("queue mine", reverse("requests:approval-list"), {"queue": "mine"}),
        ("queue hr", reverse("requests:approval-list"), {"queue": "hr"}),
        ("queue certificates", reverse("requests:approval-list"), {"queue": "certificates"}),
        ("queue stuck", reverse("requests:approval-list"), {"queue": "stuck"}),
        # My space's home lists the requester's own requests in the shared RequestOut (master 7.11),
        # which carries their own comment; whoever else calls it gets their own home.
        ("home me", reverse("reports:home-me"), {}),
    ]


def _closed_urls(world):
    """Places where the marker must never appear, whoever asks."""
    owner = {"public_id": world.owner.public_id}
    urls = [
        ("notifications", reverse("notifications:list"), {"page_size": 100}),
        ("audit business", reverse("audit:business"), {"page_size": 100}),
        ("audit business export", reverse("audit:business-export"), {}),
        ("audit security", reverse("audit:security"), {"page_size": 100}),
        ("audit security export", reverse("audit:security-export"), {}),
        ("calendar department", reverse("calendar:feed"), {**PERIOD, "scope": "department"}),
        ("calendar team", reverse("calendar:feed"), {**PERIOD, "scope": "team"}),
        ("calendar company", reverse("calendar:feed"), {**PERIOD, "scope": "company"}),
        ("home team", reverse("reports:home-team"), {}),
        ("home overview", reverse("reports:home-overview"), {}),
        ("workspace summary", reverse("reports:workspace-summary"), {}),
        ("kpmg export", reverse("reports:kpmg-export"), PERIOD),
        ("kpmg default period", reverse("reports:kpmg-default-period"), {}),
        ("report filters", reverse("reports:report-filters"), {}),
        ("employee", reverse("organisation:employee-detail", kwargs=owner), {}),
        ("employee history", reverse("organisation:employee-history", kwargs=owner), {}),
        ("employees", reverse("organisation:employee-list"), {"page_size": 100}),
        ("balance", reverse("balances:figures", kwargs={"employee_public_id": world.owner.public_id}), {}),
        ("ledger", reverse("balances:ledger", kwargs={"employee_public_id": world.owner.public_id}), {}),
        ("admin status", reverse("core:admin-status"), {}),
        ("holiday impact", reverse("calendar:holiday-impact"), {"date": "2026-11-02"}),
        ("employee me", reverse("organisation:employee-me"), {}),
        ("employee documents", reverse("organisation:employee-documents", kwargs=owner), {}),
        ("employee archive preview", reverse("organisation:employee-archive-preview", kwargs=owner), {}),
        ("directory", reverse("organisation:employee-directory"), {}),
        ("accounts", reverse("accounts:account-list"), {"page_size": 100}),
        # An account's detail carries its recent audit events (master 7.2).
        ("account", reverse("accounts:account-detail", kwargs={"public_id": world.owner_user.public_id}), {}),
        ("me", reverse("accounts:me"), {}),
        ("unread count", reverse("notifications:unread-count"), {}),
    ]
    for name, query in REPORT_QUERIES.items():
        url = reverse("reports:report", kwargs={"name": name})
        urls.append((f"report {name}", url, query))
        urls.append((f"report {name} xlsx", url, {**query, "format": "xlsx"}))
    return urls


def _build(world, stage: str):
    give_balance(world, world.owner, "20")
    request = make_request(
        world,
        user=world.owner_user,
        employee=world.owner,
        request_type=world.sick,
        start=date(2026, 11, 2),
        end=date(2026, 11, 3),
        comment=f"{MARKER} migraine with fever",
    )
    link = request_services.attach_file(
        actor=world.owner_user, request=request, upload=pdf(f"{MARKER}-certificat.pdf")
    )
    request.refresh_from_db()
    if stage in ("approved", "corrected"):
        request = decide(request, world.hr_user)
    if stage == "corrected":
        request = request_services.hr_correct(
            actor=world.hr_user,
            request=request,
            reason="Returned one day earlier",
            data=RequestData(
                start_date=date(2026, 11, 2),
                end_date=date(2026, 11, 2),
                comment=f"{MARKER} migraine with fever",
            ),
        )
    if stage == "reclassified":
        request = request_services.reclassify_as_annual(actor=world.hr_user, request=request)
    return request, link.attachment


# GET routes that hold no request, comment, file name or audit row of anybody: configuration, the reference
# lists, other people's records that a request never writes to, and the plumbing. Every other GET route must
# be read by the sweep above, so that a new reading endpoint is swept or named here (review focus 9).
NOTHING_OF_A_REQUEST = {
    "accounts:account-candidates",  # employees without an account: names and addresses
    "accounts:csrf",
    "core:health",
    "calendar:holiday-list",
    "organisation:delegation-list",
    "organisation:department-list",
    "organisation:position-list",
    "organisation:profile-change-detail",  # a person's own profile fields, never a request
    "organisation:profile-change-list",
    "requests:event-list",  # exceptional events HR declares
    "requests:type-chain-preview",
    "requests:type-detail",
    "requests:type-list",
    "requests:type-steps",
    "settings:hr",
    "settings:system",
}


@pytest.mark.django_db
def test_every_reading_endpoint_is_swept_or_holds_nothing_of_a_request(world):
    from django.urls import resolve

    from tests.hardening.routes import iter_api_routes

    request, attachment = _build(world, "pending")
    swept = {
        resolve(url).namespace + ":" + resolve(url).url_name
        for _, url, _ in _detail_urls(world, request, attachment) + _closed_urls(world)
    }
    reading = {route.name for route in iter_api_routes() if "GET" in route.methods}

    assert sorted(reading - swept - NOTHING_OF_A_REQUEST) == []
    assert sorted(NOTHING_OF_A_REQUEST - reading) == [], "a route named above no longer exists"
    assert sorted(NOTHING_OF_A_REQUEST & swept) == [], "a route named above is swept: remove it from the list"


@pytest.mark.django_db
@pytest.mark.parametrize("stage", STAGES)
def test_the_marker_reaches_only_the_requester_and_hr(world, client_for, frozen, caplog, stage):
    caplog.set_level(logging.DEBUG)
    request, attachment = _build(world, stage)
    leaks: list[str] = []
    answered: dict[str, set[int]] = {}

    for persona in EVERYONE:
        client = client_for(world.user_for(persona))
        for label, url, params in _detail_urls(world, request, attachment):
            response = client.get(url, params)
            answered.setdefault(label, set()).add(response.status_code)
            if persona not in MAY_SEE and MARKER in _text(response):
                leaks.append(f"{stage}: {persona} read the marker in '{label}' ({url})")

    # The downloads above were audited; the closed places are read last so they include those rows.
    for persona in EVERYONE:
        client = client_for(world.user_for(persona))
        for label, url, params in _closed_urls(world):
            response = client.get(url, params)
            answered.setdefault(label, set()).add(response.status_code)
            if MARKER in _text(response):
                leaks.append(f"{stage}: {persona} read the marker in '{label}' ({url})")

    assert leaks == []

    # The sweep is only meaningful if every place answered 200 to somebody.
    silent = sorted(label for label, codes in answered.items() if 200 not in codes)
    assert silent == [], f"no persona received 200 from: {silent}"


@pytest.mark.django_db
@pytest.mark.parametrize("stage", STAGES)
def test_the_requester_and_hr_still_see_their_own_content(world, client_for, frozen, stage):
    request, attachment = _build(world, stage)
    detail = reverse("requests:request-detail", kwargs={"public_id": request.public_id})
    download = reverse("files:download", kwargs={"public_id": attachment.public_id})
    for user in (world.owner_user, world.hr_user):
        body = client_for(user).get(detail).json()
        assert MARKER in (body["comment"] or ""), f"{user.email} lost the comment at stage {stage}"
        assert body["comment_hidden"] is False
        assert any(MARKER in item["name"] for item in body["attachments"])
        assert client_for(user).get(download).status_code == 200
    for user in (world.manager_user, world.exec_user):
        body = client_for(user).get(detail).json()
        assert body["comment"] is None and body["comment_hidden"] is True
        assert body["attachments"] == [] or all(item["kind"] != "medical" for item in body["attachments"])
        assert body["certificate"] is None


@pytest.mark.django_db
@pytest.mark.parametrize("stage", STAGES)
def test_the_marker_is_in_no_audit_row_no_notification_and_no_log(world, client_for, frozen, caplog, stage):
    caplog.set_level(logging.DEBUG)
    request, attachment = _build(world, stage)
    download = reverse("files:download", kwargs={"public_id": attachment.public_id})
    assert client_for(world.hr_user).get(download).status_code == 200

    for event in AuditEvent.objects.all():
        row = json.dumps([event.action, event.object_label, event.changes], default=str)
        assert MARKER not in row, f"audit row {event.action} carries the marker"
    assert AuditEvent.objects.filter(action="attachment.downloaded").exists()

    for notification in Notification.objects.all():
        row = json.dumps([notification.event_code, notification.parameters], default=str)
        assert MARKER not in row, f"notification {notification.event_code} carries the marker"

    for record in caplog.records:
        assert MARKER not in json.dumps(record.__dict__, default=str), (
            f"log record from {record.name} carries the marker"
        )


@pytest.mark.django_db
def test_a_manager_is_told_of_the_sick_leave_without_any_detail(world, client_for, frozen):
    request, _ = _build(world, "pending")
    rows = client_for(world.manager_user).get(reverse("notifications:list")).json()["results"]
    declared = [row for row in rows if row["event"] == "sick.declared"]
    assert len(declared) == 1
    assert MARKER not in declared[0]["title"] + declared[0]["body"]
    assert declared[0]["link"].endswith(str(request.public_id))
