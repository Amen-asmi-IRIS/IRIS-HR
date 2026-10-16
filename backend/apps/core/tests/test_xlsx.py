from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.core.xlsx import Column, build_workbook

COLUMNS = [
    Column("name", "Nom"),
    Column("days", "Jours", "number"),
    Column("start", "Début", "date"),
    Column("seen", "Vu le", "datetime"),
]


def sheet_of(data: bytes):
    return load_workbook(BytesIO(data)).active


def test_a_workbook_has_a_header_and_typed_rows():
    row = {
        "name": "Sami Ben Ali",
        "days": Decimal("2.5"),
        "start": date(2026, 11, 3),
        "seen": datetime(2026, 10, 6, 9, 30, tzinfo=UTC),  # 10:30 in Africa/Tunis
    }

    sheet = sheet_of(build_workbook(title="Absences", columns=COLUMNS, rows=[row]))

    assert sheet.title == "Absences"
    assert [cell.value for cell in sheet[1]] == ["Nom", "Jours", "Début", "Vu le"]
    assert sheet["A2"].value == "Sami Ben Ali"
    assert sheet["B2"].value == 2.5
    assert sheet["C2"].value == datetime(2026, 11, 3)
    assert sheet["D2"].value == datetime(2026, 10, 6, 10, 30)
    assert sheet.max_row == 2


def test_the_header_is_bold_and_stays_visible():
    sheet = sheet_of(build_workbook(title="Absences", columns=COLUMNS, rows=[]))

    assert sheet["A1"].font.bold is True
    assert sheet.freeze_panes == "A2"


def test_missing_and_empty_values_leave_the_cell_empty():
    sheet = sheet_of(build_workbook(title="T", columns=COLUMNS, rows=[{"name": None, "days": ""}]))

    assert [cell.value for cell in sheet[2]] == [None, None, None, None]


@pytest.mark.parametrize("text", ["=1+1", "+SUM(A1:A9)", "-2+3", "@cmd"])
def test_text_can_never_become_a_formula(text):
    sheet = sheet_of(build_workbook(title="T", columns=[Column("name", "Nom")], rows=[{"name": text}]))

    assert sheet["A2"].value == text
    assert sheet["A2"].data_type == "s"


def test_control_characters_are_dropped_instead_of_breaking_the_file():
    rows = [{"name": "Sami\x0b Ben\x00 Ali"}]

    sheet = sheet_of(build_workbook(title="T", columns=[Column("name", "Nom")], rows=rows))

    assert sheet["A2"].value == "Sami Ben Ali"


def test_a_number_given_as_text_is_stored_as_a_number():
    sheet = sheet_of(
        build_workbook(title="T", columns=[Column("days", "Jours", "number")], rows=[{"days": "3.00"}])
    )

    assert sheet["A2"].value == 3
    assert sheet["A2"].data_type == "n"


def test_a_value_of_the_wrong_kind_is_written_as_text():
    columns = [Column("days", "Jours", "number"), Column("start", "Début", "date")]

    sheet = sheet_of(build_workbook(title="T", columns=columns, rows=[{"days": "many", "start": "soon"}]))

    assert sheet["A2"].value == "many"
    assert sheet["B2"].value == "soon"


def test_the_sheet_title_is_made_valid():
    title = "Demandes: en/attente [2026] et un titre beaucoup trop long pour Excel"

    sheet = sheet_of(build_workbook(title=title, columns=COLUMNS, rows=[]))

    assert len(sheet.title) <= 31
    assert not set(sheet.title) & set("[]:*?/\\")


def test_an_unknown_column_kind_is_refused():
    with pytest.raises(ValueError, match="colour"):
        Column("x", "X", "colour")


def test_notes_go_to_a_sheet_of_their_own_and_stay_text():
    data = build_workbook(
        title="KPMG",
        columns=[Column("name", "Nom")],
        rows=[{"name": "Sami Ben Ali"}],
        notes=["Le solde comprend des crédits à venir.", "=1+1"],
    )

    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == ["KPMG", "Notes"]
    assert workbook.active.title == "KPMG"
    assert [row for row in workbook["KPMG"].iter_rows(values_only=True)] == [("Nom",), ("Sami Ben Ali",)]
    notes = workbook["Notes"]
    assert [notes["A1"].value, notes["A2"].value] == ["Le solde comprend des crédits à venir.", "=1+1"]
    assert notes["A2"].data_type == "s"


def test_no_notes_no_second_sheet():
    data = build_workbook(title="T", columns=[Column("name", "Nom")], rows=[])

    assert load_workbook(BytesIO(data)).sheetnames == ["T"]
