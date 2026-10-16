"""Excel workbooks for exports: one data sheet, a bold header row, typed cells; notes on a sheet of their own."""

import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

KINDS = ("text", "number", "date", "datetime")
_FORBIDDEN_IN_TITLE = re.compile(r"[\[\]:*?/\\]")
NOTES_SHEET = "Notes"


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    kind: str = "text"

    def __post_init__(self) -> None:
        if self.kind not in KINDS:
            raise ValueError(f"Unknown column kind: {self.kind}")


def _sheet_title(title: str) -> str:
    return (_FORBIDDEN_IN_TITLE.sub(" ", title).strip() or "Sheet")[:31]


def _as_number(value: Any) -> int | float | Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return value
    if isinstance(value, str):
        try:
            return Decimal(value.strip())
        except InvalidOperation:
            return None
    return None


def _write_text(cell, text: str) -> None:
    """Writes an explicit string, so "=..." can never become a formula."""
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", text)  # control characters would make the file unreadable
    cell.data_type = "s"


def _write(cell, value: Any, kind: str) -> None:
    if value is None or value == "":
        return
    if kind == "number":
        number = _as_number(value)
        if number is not None:
            cell.value = number
            return
    elif kind == "date" and isinstance(value, date):
        if isinstance(value, datetime):
            value = (timezone.localtime(value) if timezone.is_aware(value) else value).date()
        cell.value = value
        cell.number_format = "DD/MM/YYYY"
        return
    elif kind == "datetime" and isinstance(value, datetime):
        local = timezone.localtime(value) if timezone.is_aware(value) else value
        cell.value = local.replace(tzinfo=None)  # Excel has no time zones: write local time
        cell.number_format = "DD/MM/YYYY HH:MM"
        return
    # Everything else is written as an explicit string, so "=..." can never become a formula.
    _write_text(cell, str(value))


def build_workbook(
    *,
    title: str,
    columns: Sequence[Column],
    rows: Iterable[Mapping[str, Any]],
    notes: Sequence[str] = (),
) -> bytes:
    """The data sheet holds the header row and the rows, nothing else, so that a program can import it.

    Sentences the reader must have beside the figures go to a second sheet, "Notes", one per line.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(title)

    widths = []
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        _write_text(cell, column.header)
        cell.font = Font(bold=True)
        widths.append(len(column.header))
    sheet.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        for index, column in enumerate(columns, start=1):
            value = row.get(column.key)
            _write(sheet.cell(row=row_index, column=index), value, column.kind)
            if value is not None:
                widths[index - 1] = max(widths[index - 1], min(len(str(value)), 60))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width + 2

    if notes:
        notes_sheet = workbook.create_sheet(NOTES_SHEET)
        for row_index, note in enumerate(notes, start=1):
            _write_text(notes_sheet.cell(row=row_index, column=1), note)
        notes_sheet.column_dimensions["A"].width = min(max(len(note) for note in notes), 120) + 2

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
