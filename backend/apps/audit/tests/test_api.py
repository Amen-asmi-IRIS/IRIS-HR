import json
import uuid
from datetime import UTC, datetime
from io import BytesIO

import pytest
import time_machine
from openpyxl import load_workbook

from apps.audit import services as audit
from apps.audit.models import AuditEvent
from apps.core import context

pytestmark = pytest.mark.django_db

BUSINESS = "/api/v1/audit/business"
SECURITY = "/api/v1/audit/security"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EVENT_KEYS = {
    "id",
    "at",
    "actor",
    "category",
    "action",
    "object_type",
    "object_id",
    "object_label",
    "changes",
    "ip_address",
    "request_id",
}


def at(day: int) -> datetime:
    return datetime(2026, 10, day, 10, 0, tzinfo=UTC)


@pytest.fixture
def events(world):
    """Three business events (one by the system) and one security event, on different days."""
    with time_machine.travel(at(1), tick=False):
        world.updated = audit.record(
            action="employee.updated",
            object_type="employee",
            object_id="emp-1",
            object_label="Amal Employee",
            changes={"phone": {"before": "111", "after": "222"}},
            actor=world.hr_user,
        )
    with time_machine.travel(at(3), tick=False):
        world.archived = audit.record(
            action="employee.archived",
            object_type="employee",
            object_id="emp-1",
            object_label="Amal Employee",
            actor=world.hr_user,
        )
    with time_machine.travel(at(4), tick=False):
        world.credited = audit.record(action="credit.posted", object_type="ledger", object_id="7", actor=None)
    with time_machine.travel(at(5), tick=False), context.overriding(ip="10.0.0.9", request_id="rid-9"):
        world.assigned = audit.record(
            action="role.assigned",
            category=audit.SECURITY,
            obj=world.owner_user,
            changes={"role": "exec"},
            actor=world.sysadmin_user,
        )
    return world


def test_hr_reads_business_events_newest_first(events, client_for):
    body = client_for(events.hr_user).get(BUSINESS).json()

    assert set(body) == {"count", "next", "previous", "results"}
    assert body["count"] == 3
    assert [event["action"] for event in body["results"]] == [
        "credit.posted",
        "employee.archived",
        "employee.updated",
    ]
    oldest = body["results"][2]
    assert set(oldest) == EVENT_KEYS
    assert oldest["actor"] == "Senna RH"
    assert oldest["category"] == "business"
    assert (oldest["object_type"], oldest["object_id"], oldest["object_label"]) == (
        "employee",
        "emp-1",
        "Amal Employee",
    )
    assert oldest["changes"] == {"phone": {"before": "111", "after": "222"}}
    assert datetime.fromisoformat(oldest["at"]) == at(1)
    assert body["results"][0]["actor"] is None  # the system


def test_the_security_log_holds_security_events_only(events, client_for):
    body = client_for(events.sysadmin_user).get(SECURITY).json()

    assert body["count"] == 1
    event = body["results"][0]
    assert event["action"] == "role.assigned"
    assert event["actor"] == "Sonia Admin"
    assert event["object_id"] == str(events.owner_user.public_id)
    assert event["ip_address"] == "10.0.0.9"
    assert event["request_id"] == "rid-9"


def test_each_log_is_closed_to_the_other_role(events, client_for):
    assert client_for(events.hr_user).get(SECURITY).status_code == 403
    assert client_for(events.sysadmin_user).get(BUSINESS).status_code == 403
    assert client_for(events.owner_user).get(BUSINESS).status_code == 403
    assert client_for(events.exec_user).get(BUSINESS).status_code == 200
    assert client_for(events.exec_user).get(SECURITY).status_code == 200


@pytest.mark.parametrize(
    ("query", "expected"),
    [
        ({"action": "employee.archived"}, ["employee.archived"]),
        ({"object_type": "employee", "object_id": "emp-1"}, ["employee.archived", "employee.updated"]),
        ({"object_type": "ledger"}, ["credit.posted"]),
        ({"from": "2026-10-02", "to": "2026-10-03"}, ["employee.archived"]),
        ({"from": "2026-10-04"}, ["credit.posted"]),
        ({"to": "2026-10-01"}, ["employee.updated"]),
        ({"action": "nothing.like.this"}, []),
    ],
)
def test_the_list_can_be_filtered(events, client_for, query, expected):
    body = client_for(events.hr_user).get(BUSINESS, query).json()

    assert [event["action"] for event in body["results"]] == expected


def test_the_list_can_be_filtered_by_actor(events, client_for):
    client = client_for(events.hr_user)

    by_hr = client.get(BUSINESS, {"actor": str(events.hr_user.public_id)}).json()
    by_nobody = client.get(BUSINESS, {"actor": str(uuid.uuid4())}).json()

    assert [event["action"] for event in by_hr["results"]] == ["employee.archived", "employee.updated"]
    assert by_nobody["count"] == 0


@pytest.mark.parametrize(
    ("query", "code", "field"),
    [
        ({"from": "2026-10-05", "to": "2026-10-01"}, "audit.period_order", "to"),
        ({"from": "yesterday"}, "core.validation", "from"),
        ({"actor": "not-a-uuid"}, "core.validation", "actor"),
    ],
)
def test_a_wrong_filter_is_explained(events, client_for, query, code, field):
    response = client_for(events.hr_user).get(BUSINESS, query)

    assert response.status_code == 400
    assert response.json()["code"] == code
    assert response.json()["errors"][0]["field"] == field


def test_the_filtered_list_can_be_exported_to_excel(events, client_for):
    events.hr_user.language = "en"
    events.hr_user.save(update_fields=["language"])

    response = client_for(events.hr_user).get(f"{BUSINESS}/export", {"object_type": "employee"})

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX
    disposition = response["Content-Disposition"]
    assert disposition.startswith('attachment; filename="business-log-')
    assert disposition.endswith('.xlsx"')
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "Activity log"
    assert [cell.value for cell in sheet[1]] == [
        "When",
        "Who",
        "Action",
        "Kind of item",
        "Item",
        "Changes",
        "Network address",
        "Reference",
    ]
    assert sheet.max_row == 3
    assert (sheet["B2"].value, sheet["C2"].value, sheet["E2"].value) == (
        "Senna RH",
        "employee.archived",
        "Amal Employee",
    )
    assert sheet["A3"].value == datetime(2026, 10, 1, 11, 0)  # 10:00 UTC, written in Africa/Tunis time
    assert json.loads(sheet["F3"].value) == {"phone": {"before": "111", "after": "222"}}


def test_an_export_is_itself_audited_with_its_filters(events, client_for):
    client_for(events.hr_user).get(f"{BUSINESS}/export", {"object_type": "employee", "from": "2026-10-01"})
    client_for(events.sysadmin_user).get(f"{SECURITY}/export")

    business = AuditEvent.objects.get(action="audit.exported", category="business")
    security = AuditEvent.objects.get(action="audit.exported", category="security")
    assert business.actor == events.hr_user
    assert business.changes == {"filters": {"object_type": "employee", "from": "2026-10-01"}, "rows": 2}
    assert security.actor == events.sysadmin_user
    assert security.changes == {"filters": {}, "rows": 1}


def test_the_export_is_written_in_the_readers_language_and_names_the_system(events, client_for):
    events.hr_user.language = "fr"
    events.hr_user.save(update_fields=["language"])

    response = client_for(events.hr_user).get(f"{BUSINESS}/export", {"object_type": "ledger"})

    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "Journal d'activité"
    assert sheet["A1"].value == "Date et heure"
    assert sheet["B2"].value == "Système"


def test_exports_follow_the_same_access_as_the_lists(events, client_for, anon_client):
    assert client_for(events.hr_user).get(f"{SECURITY}/export").status_code == 403
    assert client_for(events.sysadmin_user).get(f"{BUSINESS}/export").status_code == 403
    assert anon_client.get(f"{BUSINESS}/export").status_code == 401
