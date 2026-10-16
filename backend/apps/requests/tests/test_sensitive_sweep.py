"""Master section 10, item 9: nothing medical reaches a manager, an executive, a colleague or the trail."""

import json
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.notifications.models import Notification
from apps.reports.tabular.registry import REPORTS
from apps.reports.tabular.render import XLSX_MEDIA_TYPE
from apps.requests import services
from apps.requests.tests.helpers import (
    MON,
    NOW,
    TODAY,
    TUE,
    days,
    decide,
    fresh,
    pdf,
    pending_step,
    submitted,
)

pytestmark = pytest.mark.django_db

SECRETS = ("zz-secret-grippe", "zz-secret-certificat", "zz-secret-retour")


@pytest.fixture
def story(world):
    """A sick leave with a comment, a certificate, a return by HR with a comment, and a resubmission."""
    request = submitted(world, request_type=world.sick, data=days(TODAY, comment="ZZ-SECRET-GRIPPE sévère"))
    services.attach_file(actor=world.owner_user, request=request, upload=pdf("zz-secret-certificat.pdf"))
    decide(request, world.hr_user, "return", "ZZ-SECRET-RETOUR : certificat illisible")
    returned = fresh(request)
    services.submit(actor=world.owner_user, request=returned, expected_version=returned.version)
    return request


def everything_readable_by(client, request) -> str:
    """The body of every read endpoint that could mention the request, whatever the status of the answer."""
    one = {"public_id": request.public_id}
    person = {"employee_public_id": request.employee.public_id}
    window = {"from": request.start_date.isoformat(), "to": request.end_date.isoformat()}
    calls = [
        (reverse("requests:request-list"), {}),
        (reverse("requests:request-list"), {"scope": "team"}),
        (reverse("requests:request-list"), {"scope": "all"}),
        (reverse("requests:request-list"), {"scope": "decided"}),
        (reverse("requests:request-detail", kwargs=one), {}),
        (reverse("requests:request-history", kwargs=one), {}),
        (reverse("requests:request-attachments", kwargs=one), {}),
        (reverse("requests:request-amendments", kwargs=one), {}),
        (reverse("requests:approval-list"), {}),
        (reverse("requests:approval-list"), {"queue": "hr"}),
        (reverse("requests:approval-list"), {"queue": "certificates"}),
        (reverse("requests:approval-list"), {"queue": "stuck"}),
        (reverse("notifications:list"), {}),
        (reverse("audit:business"), {}),
        (reverse("balances:figures", kwargs=person), {}),
        (reverse("balances:ledger", kwargs=person), {}),
        *(
            (reverse("calendar:feed"), {**window, "scope": scope})
            for scope in ("department", "team", "company")
        ),
        (reverse("reports:home-me"), {}),
        (reverse("reports:home-team"), {}),
        (reverse("reports:home-overview"), {}),
        (reverse("reports:report-filters"), {}),
        *(
            (reverse("reports:report", kwargs={"name": name}), {**window, **as_file})
            for name in REPORTS
            for as_file in ({}, {"format": "xlsx"})
        ),
        (reverse("reports:kpmg-export"), window),
    ]
    return "\n".join(_text(client.get(url, params)) for url, params in calls).lower()


def _text(response) -> str:
    """The body of an answer as text; a workbook (a report, the KPMG file) as the text of the cells of every sheet."""
    if response["Content-Type"] != XLSX_MEDIA_TYPE:
        return response.content.decode()
    workbook = load_workbook(BytesIO(response.content))
    rows = (row for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
    return " ".join(str(cell) for row in rows for cell in row if cell is not None)


def test_no_endpoint_leaks_the_sick_leave_comment(world, client_for, story):
    others = (
        world.manager_user,
        world.exec_user,
        world.peer_user,
        world.other_manager_user,
        world.outsider_user,
        world.sysadmin_user,
    )
    for user in others:
        text = everything_readable_by(client_for(user), story)
        for secret in SECRETS:
            assert secret not in text, f"{secret} reached {user.email}"


def test_the_manager_still_sees_that_the_sick_leave_exists(world, client_for, story):
    body = (
        client_for(world.manager_user)
        .get(reverse("requests:request-detail", kwargs={"public_id": story.public_id}))
        .json()
    )
    assert (body["type"]["code"], body["status"], body["comment"], body["comment_hidden"]) == (
        "sick",
        "pending",
        None,
        True,
    )
    assert (body["certificate"], body["attachments"]) == (None, [])


def test_the_requester_and_hr_read_all_of_it(world, client_for, story):
    for user in (world.owner_user, world.hr_user):
        text = (
            client_for(user)
            .get(reverse("requests:request-detail", kwargs={"public_id": story.public_id}))
            .content.decode()
            .lower()
        )
        for secret in SECRETS:
            assert secret in text, f"{secret} is missing for {user.email}"


def test_nothing_medical_is_stored_in_notifications_or_in_the_audit_trail(world, story):
    stored = json.dumps(
        [
            list(Notification.objects.values_list("event_code", "parameters", "target_public_id")),
            list(AuditEvent.objects.values_list("action", "changes", "object_label")),
        ],
        default=str,
    ).lower()
    for secret in SECRETS:
        assert secret not in stored
    assert AuditEvent.objects.filter(action="step.decided").exists()
    assert Notification.objects.filter(event_code="request.returned").exists()


# ------------------------------------------------------------------ what is written after the declaration

LATER_SECRETS = (
    "zz-secret-grippe",
    "zz-secret-exception",
    "zz-secret-raison",
    "zz-secret-note",
    "zz-secret-scan",
    "zz-secret-correction",
    "zz-secret-changement",
    "zz-secret-annulation",
)


@pytest.fixture
def long_story(world):
    """A sick leave HR entered with an exception, then every action that carries free text or a file.

    The step is given to someone outside the team, their decision is recorded from paper with its scan,
    HR corrects the dates, the employee asks to cancel, and HR cancels.
    """
    request = submitted(
        world,
        actor=world.hr_user,
        request_type=world.sick,
        data=days(MON, TUE, comment="ZZ-SECRET-GRIPPE sévère"),
        override_reason="ZZ-SECRET-EXCEPTION : hospitalisation",
    )
    services.reassign_step(
        actor=world.hr_user,
        step=pending_step(request),
        new_approver=world.other_manager,
        reason="ZZ-SECRET-RAISON : le médecin du travail tranche",
    )
    current = fresh(request)
    services.record_on_behalf(
        actor=world.hr_user,
        step=pending_step(current),
        decided_by=world.other_manager,
        decision="approve",
        decided_at=NOW,
        channel="paper",
        note="ZZ-SECRET-NOTE : accord après lecture du dossier",
        scan=pdf("zz-secret-scan.pdf"),
        expected_version=current.version,
    )
    services.hr_correct(
        actor=world.hr_user,
        request=fresh(request),
        data=days(MON),
        reason="ZZ-SECRET-CORRECTION : sortie anticipée",
    )
    services.open_amendment(
        actor=world.owner_user,
        request=fresh(request),
        kind="cancel",
        data=None,
        reason="ZZ-SECRET-CHANGEMENT : rétabli",
    )
    services.hr_cancel(actor=world.hr_user, request=fresh(request), reason="ZZ-SECRET-ANNULATION : doublon")
    return fresh(request)


def test_no_endpoint_leaks_what_was_written_on_a_sick_leave_afterwards(world, client_for, long_story):
    assert long_story.status == "cancelled"
    others = (
        world.manager_user,
        world.exec_user,
        world.peer_user,
        world.other_manager_user,  # was given the step and is recorded as having decided it
        world.outsider_user,
        world.sysadmin_user,
    )
    for user in others:
        text = everything_readable_by(client_for(user), long_story)
        for secret in LATER_SECRETS:
            assert secret not in text, f"{secret} reached {user.email}"


def test_the_one_who_decided_a_sick_leave_from_outside_the_team_sees_it_without_its_content(
    world, client_for, long_story
):
    body = client_for(world.other_manager_user).get(
        reverse("requests:request-detail", kwargs={"public_id": long_story.public_id})
    )
    assert body.status_code == 200
    seen = body.json()
    assert (seen["comment"], seen["comment_hidden"], seen["certificate"]) == (None, True, None)
    assert (seen["attachments"], seen["decision_context"]) == ([], None)
    assert all(step["comment"] is None for step in seen["steps"])
    assert all(amendment["reason"] is None for amendment in seen["amendments"])


def test_nothing_written_afterwards_is_stored_in_notifications_or_in_the_audit_trail(world, long_story):
    stored = json.dumps(
        [
            list(Notification.objects.values_list("event_code", "parameters", "target_public_id")),
            list(AuditEvent.objects.values_list("action", "changes", "object_label")),
        ],
        default=str,
    ).lower()
    for secret in LATER_SECRETS:
        assert secret not in stored
    actions = set(AuditEvent.objects.values_list("action", flat=True))
    assert {
        "request.created",
        "step.reassigned",
        "step.recorded_on_behalf",
        "request.corrected",
        "amendment.opened",
        "request.cancelled",
    } <= actions


def test_the_requester_and_hr_still_read_what_is_kept_on_the_request(world, client_for, long_story):
    """The comment, the note of the recorded decision, the reason of the change and the scan stay readable."""
    for user in (world.owner_user, world.hr_user):
        text = (
            client_for(user)
            .get(reverse("requests:request-detail", kwargs={"public_id": long_story.public_id}))
            .content.decode()
            .lower()
        )
        for secret in ("zz-secret-grippe", "zz-secret-note", "zz-secret-scan", "zz-secret-changement"):
            assert secret in text, f"{secret} is missing for {user.email}"
