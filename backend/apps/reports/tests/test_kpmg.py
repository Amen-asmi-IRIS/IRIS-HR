from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import translation
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.balances import services as balance_services
from apps.core.errors import DomainError
from apps.organisation import selectors as organisation_selectors
from apps.organisation import services as organisation_services
from apps.reports.kpmg import KPMG_COLUMNS, build_kpmg, default_period, export_kpmg, months_covered
from apps.reports.tests import scenario
from apps.requests import services as request_services
from apps.settings import services as settings_service

pytestmark = pytest.mark.django_db

NOVEMBER = {"from": "2026-11-01", "to": "2026-11-30"}


def _rows(user, query=None, **extra):
    with translation.override("fr"):
        _, result = build_kpmg(user=user, query={**(query or NOVEMBER), **extra})
    return {row["employee_number"]: row for row in result.rows}


def _grid(content: bytes) -> list[tuple]:
    return list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))


def test_the_ten_columns_of_the_file(world, november):
    with translation.override("fr"):
        _, result = build_kpmg(user=world.hr_user, query=NOVEMBER)

    assert [column.key for column in result.columns] == [key for key, _ in KPMG_COLUMNS]
    assert [column.label for column in result.columns] == [
        "Matricule",
        "Nom",
        "Département",
        "Congé annuel (jours)",
        "Congé maladie (jours)",
        "Congé exceptionnel (jours)",
        "Permissions (heures)",
        "Télétravail (jours)",
        "Solde annuel en fin de période",
        "Tickets restaurant",
    ]
    assert result.totals is None


def test_november_matches_the_hand_calculation(world, november):
    rows = _rows(world.hr_user)

    owner = rows[world.owner.employee_number]
    assert owner["name"] == organisation_selectors.full_name(world.owner)
    assert owner["annual_days"] == Decimal("1.5")
    assert owner["sick_days"] == Decimal("0.5")
    assert owner["exceptional_days"] == Decimal("1")
    assert owner["permission_hours"] == Decimal("2.5")
    assert owner["remote_days"] == Decimal("2")
    assert owner["balance_end"] == Decimal("16.5")
    assert owner["meal_tickets"] == 18

    outsider = rows[world.outsider.employee_number]
    assert outsider["sick_days"] == Decimal("3")
    assert outsider["annual_days"] == Decimal("0")
    assert outsider["balance_end"] == Decimal("0")
    assert outsider["meal_tickets"] == 18

    peer = rows[world.peer.employee_number]
    assert [peer[key] for key, kind in KPMG_COLUMNS[3:8]] == [Decimal("0")] * 5
    assert peer["meal_tickets"] == 21


def test_one_row_per_employee(world, november):
    rows = _rows(world.hr_user)

    assert set(rows) == {
        employee.employee_number
        for employee in (
            world.ceo,
            world.hr,
            world.manager,
            world.owner,
            world.peer,
            world.other_manager,
            world.outsider,
        )
    }


def test_only_approved_requests_count(world, november):
    before = _rows(world.hr_user)[world.owner.employee_number]
    assert before["annual_days"] == Decimal("1.5")

    scenario.approve_fully(november.pending_annual)  # 9 and 10 November

    after = _rows(world.hr_user)[world.owner.employee_number]
    assert after["annual_days"] == Decimal("3.5")
    assert after["meal_tickets"] == 16


def test_two_half_days_on_one_day_cost_one_ticket(world, november):
    rows = _rows(world.hr_user, {"from": "2026-11-13", "to": "2026-11-13"})

    owner = rows[world.owner.employee_number]
    assert (owner["annual_days"], owner["sick_days"]) == (Decimal("0.5"), Decimal("0.5"))
    assert owner["meal_tickets"] == 20


def test_a_single_half_day_and_remote_days_cost_nothing(world, november):
    scenario.approved(
        world, employee=world.peer, request_type=world.annual, start=date(2026, 11, 16), ends_at_midday=True
    )
    scenario.approved(world, employee=world.peer, request_type=world.remote, start=date(2026, 11, 17))

    peer = _rows(world.hr_user)[world.peer.employee_number]

    assert peer["annual_days"] == Decimal("0.5")
    assert peer["remote_days"] == Decimal("1")
    assert peer["meal_tickets"] == 21


def test_type_filter_narrows_the_absence_columns_only(world, november):
    owner = _rows(world.hr_user, type="sick")[world.owner.employee_number]

    assert owner["sick_days"] == Decimal("0.5")
    assert owner["annual_days"] == Decimal("0")
    assert owner["remote_days"] == Decimal("0")
    assert owner["meal_tickets"] == 18
    assert owner["balance_end"] == Decimal("16.5")


def test_department_and_employee_filters(world, november):
    by_department = _rows(world.hr_user, department=str(world.dept_b.id))
    by_employee = _rows(world.hr_user, employee=str(world.owner.public_id))

    assert world.outsider.employee_number in by_department
    assert world.owner.employee_number not in by_department
    assert set(by_employee) == {world.owner.employee_number}


def test_nothing_to_export(world, november):
    with pytest.raises(DomainError) as refused:
        _rows(world.hr_user, department="999999")

    assert refused.value.code == "reports.nothing_to_export"


def test_someone_hired_after_the_period_is_left_out(world, november):
    organisation_services.update_employee(
        actor=world.hr_user, employee=world.peer, data={"hire_date": date(2026, 12, 1)}
    )

    assert world.peer.employee_number not in _rows(world.hr_user)


def test_a_leaver_appears_until_the_period_of_the_leaving_date(world, november):
    organisation_services.archive_employee(
        actor=world.hr_user, employee=world.peer, leaving_date=date(2026, 11, 20), reason="Resignation"
    )

    assert world.peer.employee_number in _rows(world.hr_user)  # left during November
    assert world.peer.employee_number in _rows(world.hr_user, {"from": "2026-11-20", "to": "2026-11-20"})
    assert world.peer.employee_number not in _rows(world.hr_user, {"from": "2026-11-21", "to": "2026-11-30"})
    assert world.peer.employee_number not in _rows(world.hr_user, {"from": "2026-12-01", "to": "2026-12-31"})


@pytest.mark.parametrize(
    ("start", "end", "months"),
    [
        (date(2026, 11, 1), date(2026, 11, 30), 1),
        (date(2026, 10, 26), date(2026, 11, 25), 1),
        (date(2026, 11, 10), date(2026, 11, 20), 1),
        (date(2026, 11, 30), date(2026, 12, 1), 1),
        (date(2026, 11, 1), date(2026, 12, 15), 1),
        (date(2026, 11, 1), date(2026, 12, 30), 1),
        (date(2026, 11, 1), date(2026, 12, 31), 2),
        (date(2027, 1, 29), date(2027, 2, 28), 1),
        (date(2027, 2, 2), date(2027, 3, 1), 1),
        (date(2026, 11, 29), date(2026, 12, 28), 1),
        (date(2026, 10, 26), date(2026, 12, 25), 2),
        (date(2026, 1, 1), date(2026, 12, 31), 12),
    ],
)
def test_months_covered(start, end, months):
    assert months_covered(start, end) == months


def test_a_two_month_period_doubles_the_base(world, november):
    peer = _rows(world.hr_user, {"from": "2026-11-01", "to": "2026-12-31"})[world.peer.employee_number]

    assert peer["meal_tickets"] == 42


def test_a_part_month_adds_no_base(world, november):
    """D38: the base is multiplied by whole months only; 1 November to 15 December is one month."""
    peer = _rows(world.hr_user, {"from": "2026-11-01", "to": "2026-12-15"})[world.peer.employee_number]

    assert peer["meal_tickets"] == 21


def test_the_base_is_a_setting_and_tickets_never_go_below_zero(world, november):
    settings_service.set_value(actor=world.hr_user, key="meal_tickets_per_month", value=20)
    assert _rows(world.hr_user)[world.peer.employee_number]["meal_tickets"] == 20

    settings_service.set_value(actor=world.hr_user, key="meal_tickets_per_month", value=1)
    assert _rows(world.hr_user)[world.owner.employee_number]["meal_tickets"] == 0


def test_default_period_is_the_last_month_that_is_over(world):
    """D53: with no cut-off day, the month just ended, never the running one."""
    settings_service.clear_cache()

    assert default_period(today=date(2026, 11, 18)) == (date(2026, 10, 1), date(2026, 10, 31), None)
    assert default_period(today=date(2026, 12, 3)) == (date(2026, 11, 1), date(2026, 11, 30), None)
    assert default_period(today=date(2026, 12, 31)) == (date(2026, 11, 1), date(2026, 11, 30), None)
    assert default_period(today=date(2027, 1, 1)) == (date(2026, 12, 1), date(2026, 12, 31), None)
    assert default_period(today=date(2027, 3, 10)) == (date(2027, 2, 1), date(2027, 2, 28), None)


def test_default_period_with_a_cutoff_day(world):
    settings_service.set_value(actor=world.hr_user, key="export_cutoff_day", value=25)

    assert default_period(today=date(2026, 11, 18)) == (date(2026, 10, 26), date(2026, 11, 25), 25)
    assert default_period(today=date(2026, 11, 27)) == (date(2026, 10, 26), date(2026, 11, 25), 25)
    assert default_period(today=date(2027, 3, 3)) == (date(2027, 2, 26), date(2027, 3, 25), 25)


def test_default_period_in_january(world):
    settings_service.set_value(actor=world.hr_user, key="export_cutoff_day", value=25)

    assert default_period(today=date(2027, 1, 10)) == (date(2026, 12, 26), date(2027, 1, 25), 25)


def test_default_period_with_the_first_and_the_last_cutoff_days(world):
    settings_service.set_value(actor=world.hr_user, key="export_cutoff_day", value=1)
    assert default_period(today=date(2026, 11, 18)) == (date(2026, 10, 2), date(2026, 11, 1), 1)
    assert default_period(today=date(2027, 3, 1)) == (date(2027, 2, 2), date(2027, 3, 1), 1)

    settings_service.set_value(actor=world.hr_user, key="export_cutoff_day", value=28)
    # The previous cut-off is 28 February: the period starts on 1 March.
    assert default_period(today=date(2027, 3, 3)) == (date(2027, 3, 1), date(2027, 3, 28), 28)
    assert default_period(today=date(2026, 12, 29)) == (date(2026, 11, 29), date(2026, 12, 28), 28)


def test_types_hr_adds_feed_the_columns_their_unit_and_label_say(world, november):
    """D37: every hour-based type adds to permission hours, every "remote" type to remote days; another
    day-based type feeds no column, but costs meal tickets when HR flags it so (REP-06)."""
    names = {"name_fr": "Type de test", "name_en": "Test type"}
    training = request_services.create_request_type(
        actor=world.hr_user, data={**names, "code": "training", "unit": "hour", "calendar_label": "hidden"}
    )
    site = request_services.create_request_type(
        actor=world.hr_user,
        data={**names, "code": "client_site", "unit": "day", "calendar_label": "remote"},
    )
    unpaid = request_services.create_request_type(
        actor=world.hr_user,
        data={
            **names,
            "code": "unpaid",
            "unit": "day",
            "calendar_label": "absent",
            "reduces_meal_tickets": True,
        },
    )
    peer = world.peer
    scenario.approved(
        world,
        employee=peer,
        request_type=training,
        start=date(2026, 11, 5),
        start_time=time(14, 0),
        end_time=time(15, 30),
    )
    scenario.approved(
        world, employee=peer, request_type=site, start=date(2026, 11, 16), end=date(2026, 11, 17)
    )
    scenario.approved(
        world, employee=peer, request_type=unpaid, start=date(2026, 11, 23), end=date(2026, 11, 24)
    )

    row = _rows(world.hr_user)[peer.employee_number]

    assert row["permission_hours"] == Decimal("1.5")
    assert row["remote_days"] == Decimal("2")
    assert (row["annual_days"], row["sick_days"], row["exceptional_days"]) == (0, 0, 0)
    assert row["meal_tickets"] == 19


def test_export_file_name_and_cells(world, november):
    with translation.override("fr"):
        content, filename = export_kpmg(user=world.hr_user, query=NOVEMBER)

    assert filename == "kpmg_2026-11-01_2026-11-30.xlsx"
    grid = _grid(content)
    header = next(index for index, row in enumerate(grid) if row and row[0] == "Matricule")
    owner = next(row for row in grid[header + 1 :] if row[0] == world.owner.employee_number)
    assert owner[3:10] == (1.5, 0.5, 1, 2.5, 2, 16.5, 18)


def test_export_headers_follow_the_readers_language(world, november):
    with translation.override("en"):
        content, _ = export_kpmg(user=world.hr_user, query=NOVEMBER)

    assert any(row and row[0] == "Employee number" and row[9] == "Meal tickets" for row in _grid(content))


def test_export_contains_no_comment(world, november):
    content, _ = export_kpmg(user=world.hr_user, query=NOVEMBER)

    text = " ".join(str(cell) for row in _grid(content) for cell in row if cell is not None)
    assert scenario.SICK_COMMENT not in text
    assert "Karray" not in text


def test_export_audit_entry_holds_filters_only(world, november):
    export_kpmg(user=world.hr_user, query={**NOVEMBER, "department": str(world.dept_a.id)})

    event = AuditEvent.objects.filter(action="export.kpmg").latest("at")
    assert event.category == "business"
    assert event.actor == world.hr_user
    assert (event.object_type, event.object_id) == ("export", "kpmg")
    assert event.changes["filters"] == {
        "scope": "all",
        "from": "2026-11-01",
        "to": "2026-11-30",
        "department": str(world.dept_a.id),
    }
    assert set(event.changes) == {"filters", "rows"}
    assert "Karray" not in repr(event.changes)


# --- a period that ends today or later: the month-end credits still to come (master 13.1) -----------


@pytest.fixture
def credits_from_october(world, november):
    """Monthly credits start on 1 October; the job has not run, so none is posted."""
    settings_service.set_value(actor=world.hr_user, key="credit_start_date", value="2026-10-01")
    return november


def _note_lines(content: bytes) -> list[str]:
    """The sentences of the sheet "Notes"; the data sheet never carries them (KPMG imports it)."""
    workbook = load_workbook(BytesIO(content))
    if "Notes" not in workbook.sheetnames:
        return []
    return [row[0] for row in workbook["Notes"].iter_rows(values_only=True)]


def _assert_strictly_tabular(content: bytes, result) -> None:
    """The first sheet is the header row and one row per employee, nothing before or after."""
    grid = _grid(content)
    assert grid[0][0] == "Matricule"
    assert len(grid) == 1 + len(result.rows)
    assert [row[0] for row in grid[1:]] == [row["employee_number"] for row in result.rows]


def test_a_period_already_over_counts_the_credits_not_posted_yet_and_says_so(world, credits_from_october):
    """Exported on 1 December before the month-end job ran: October and November are due."""
    with translation.override("fr"):
        _, result = build_kpmg(user=world.hr_user, query=NOVEMBER)
        content, _ = export_kpmg(user=world.hr_user, query=NOVEMBER)

    owner = next(row for row in result.rows if row["employee_number"] == world.owner.employee_number)
    assert owner["balance_end"] == Decimal("21.5")
    note = "Le solde annuel au 30/11/2026 comprend les crédits de fin de mois qui ne sont pas encore enregistrés."
    assert result.notes == (note,)
    assert _note_lines(content) == [note]
    _assert_strictly_tabular(content, result)


def test_a_period_already_over_whose_credits_are_posted_says_nothing(world, credits_from_october):
    balance_services.post_month_end_credits(today=date(2026, 12, 1))

    with translation.override("fr"):
        _, result = build_kpmg(user=world.hr_user, query=NOVEMBER)
        content, _ = export_kpmg(user=world.hr_user, query=NOVEMBER)

    owner = next(row for row in result.rows if row["employee_number"] == world.owner.employee_number)
    assert owner["balance_end"] == Decimal("21.5")
    assert result.notes == ()
    assert _note_lines(content) == []
    assert load_workbook(BytesIO(content)).sheetnames == ["KPMG 2026-11-01 2026-11-30"]
    _assert_strictly_tabular(content, result)


def test_a_period_ending_today_counts_the_credits_not_posted_yet_and_says_so(world, credits_from_october):
    query = {"from": "2026-11-02", "to": "2026-12-01"}
    with translation.override("fr"):
        _, result = build_kpmg(user=world.hr_user, query=query)
        content, _ = export_kpmg(user=world.hr_user, query=query)

    owner = next(row for row in result.rows if row["employee_number"] == world.owner.employee_number)
    # October and November, 2.5 days each, are due and not posted yet.
    assert owner["balance_end"] == Decimal("21.5")
    note = "Le solde annuel au 01/12/2026 comprend les crédits de fin de mois qui ne sont pas encore enregistrés."
    assert result.notes == (note,)
    assert _note_lines(content) == [note]
    _assert_strictly_tabular(content, result)


def test_a_period_ending_later_counts_every_credit_up_to_its_end(world, credits_from_october):
    with translation.override("en"):
        _, result = build_kpmg(user=world.hr_user, query={"from": "2026-12-01", "to": "2026-12-31"})

    owner = next(row for row in result.rows if row["employee_number"] == world.owner.employee_number)
    assert owner["balance_end"] == Decimal("24")
    assert result.notes == (
        "The annual balance on 31/12/2026 includes the month-end credits that are not posted yet.",
    )
