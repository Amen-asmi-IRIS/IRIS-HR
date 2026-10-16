"""Every endpoint of plans 05 and 06, called as every persona: nothing medical leaks, and the
department calendar never tells a colleague why somebody is away."""

import json
from datetime import date, time
from io import BytesIO

import pytest
import time_machine
from django.urls import reverse
from openpyxl import load_workbook

from apps.reports.tabular.registry import REPORTS
from apps.reports.tabular.render import XLSX_MEDIA_TYPE
from tests.matrix import P
from tests.request_flows import NOW, approved, send

pytestmark = pytest.mark.django_db

NOVEMBER = {"from": "2026-11-01", "to": "2026-11-30"}
ENDPOINTS = [
    ("calendar:feed", {}, {**NOVEMBER, "scope": "department"}),
    ("calendar:feed", {}, {**NOVEMBER, "scope": "team"}),
    ("calendar:feed", {}, {**NOVEMBER, "scope": "company"}),
    ("reports:home-me", {}, {}),
    ("reports:home-team", {}, {}),
    ("reports:home-overview", {}, {}),
    ("reports:workspace-summary", {}, {}),
    # Plan 06: the filter options, every report as a page and as a workbook, the KPMG file.
    ("reports:report-filters", {}, {}),
    *(
        ("reports:report", {"name": name}, {**NOVEMBER, "year": "2026", **as_file})
        for name in REPORTS
        for as_file in ({}, {"format": "xlsx"})
    ),
    ("reports:kpmg-export", {}, NOVEMBER),
]
SIGNED_IN = [persona for persona in P if persona is not P.ANON]
WITH_RECORD = [persona for persona in SIGNED_IN if persona is not P.SYSADMIN]


@pytest.fixture(autouse=True)
def frozen_clock():
    with time_machine.travel(NOW, tick=False):
        yield


@pytest.fixture
def scenario(world):
    """The peer (department A) and the outsider (department B) are on sick leave with a telling comment."""
    approved(
        world,
        employee=world.peer,
        user=world.peer_user,
        request_type=world.sick,
        start=date(2026, 11, 2),
        end=date(2026, 11, 3),
        comment="Severe migraine",
    )
    send(
        world,
        employee=world.peer,
        user=world.peer_user,
        request_type=world.annual,
        start=date(2026, 11, 17),
        end=date(2026, 11, 18),
        comment="Family trip",
    )
    approved(
        world,
        employee=world.peer,
        user=world.peer_user,
        request_type=world.permission,
        start=date(2026, 11, 4),
        end=date(2026, 11, 4),
        start_time=time(9, 0),
        end_time=time(10, 0),
    )
    approved(
        world,
        employee=world.outsider,
        user=world.outsider_user,
        request_type=world.sick,
        start=date(2026, 11, 2),
        end=date(2026, 11, 2),
        comment="Migraine again",
    )
    # Waiting for HR: counted as away on the homes (master D52), its comment still private.
    send(
        world,
        employee=world.outsider,
        user=world.outsider_user,
        request_type=world.sick,
        start=date(2026, 11, 4),
        end=date(2026, 11, 4),
        comment="Migraine, third day",
    )
    return world


@pytest.mark.parametrize("persona", SIGNED_IN)
def test_no_endpoint_of_this_plan_leaks_a_sick_leave_comment(scenario, client_for, persona):
    world = scenario
    client = client_for(world.user_for(persona))
    answered = 0
    for route, kwargs, query in ENDPOINTS:
        response = client.get(reverse(route, kwargs=kwargs), query)
        assert response.status_code in (200, 403), (route, query, response.status_code)
        if response.status_code == 403:
            continue
        answered += 1
        requester_reading_own_home = route == "reports:home-me" and persona in (P.PEER, P.OUTSIDER)
        if not requester_reading_own_home:
            assert "migraine" not in _text(response).lower(), (route, kwargs, query)
    assert answered >= 1  # everybody may at least read the summary


def _text(response) -> str:
    """The body of an answer as text; a workbook as the text of the cells of every sheet."""
    if response["Content-Type"] != XLSX_MEDIA_TYPE:
        return response.content.decode()
    workbook = load_workbook(BytesIO(response.content))
    rows = (row for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
    return " ".join(str(cell) for row in rows for cell in row if cell is not None)


@pytest.mark.parametrize("persona", WITH_RECORD)
def test_department_scope_never_reveals_a_colleagues_type(scenario, client_for, persona):
    world = scenario
    response = client_for(world.user_for(persona)).get(
        reverse("calendar:feed"), {**NOVEMBER, "scope": "department"}
    )

    assert response.status_code == 200
    others = [entry for entry in response.json()["entries"] if not entry["mine"]]
    for entry in others:
        assert entry["label"] in ("absent", "remote")
        assert entry["type_code"] is None
        assert entry["type_name"] is None
        assert entry["status"] is None
        assert entry["request_public_id"] is None
    text = json.dumps(others, ensure_ascii=False)
    assert world.sick.name_fr not in text and world.sick.name_en not in text
    assert world.permission.name_fr not in text
    # The colleague's pending annual leave (17 and 18 November) is not there at all.
    assert all(entry["start_date"] != "2026-11-17" for entry in others)


def test_colleagues_in_department_a_do_see_the_peer_as_absent(scenario, client_for):
    world = scenario
    body = (
        client_for(world.owner_user).get(reverse("calendar:feed"), {**NOVEMBER, "scope": "department"}).json()
    )

    assert [(entry["label"], entry["start_date"], entry["end_date"]) for entry in body["entries"]] == [
        ("absent", "2026-11-02", "2026-11-03"),
    ]


def test_a_manager_reads_the_type_of_a_sick_leave_but_never_its_comment(scenario, client_for):
    world = scenario
    response = client_for(world.manager_user).get(reverse("reports:home-team"))

    out = response.json()["who_is_out"]
    assert [person["type_name"] for person in out] == [world.sick.name_fr]
    assert "migraine" not in response.content.decode().lower()


def test_the_reports_of_hr_and_executives_reach_the_sick_leaves_without_their_comment(scenario, client_for):
    """The sweep above passes for a reason: the reports do list the two sick leaves of November."""
    world = scenario
    for user in (world.hr_user, world.exec_user):
        body = client_for(user).get(reverse("reports:report", kwargs={"name": "absences"}), NOVEMBER).json()
        sick = [row for row in body["rows"] if row["type"] in (world.sick.name_fr, world.sick.name_en)]
        assert len(sick) == 2
        assert "migraine" not in json.dumps(body).lower()
