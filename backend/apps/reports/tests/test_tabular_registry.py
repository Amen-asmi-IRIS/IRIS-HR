from io import BytesIO

import pytest
from django.utils import translation
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.core.errors import NotAllowed, NotFound
from apps.organisation import selectors as organisation_selectors
from apps.reports.tabular.registry import REPORTS, export_report, filter_options, run_report

pytestmark = pytest.mark.django_db

NOVEMBER = {"from": "2026-11-01", "to": "2026-11-30"}


def _grid(content: bytes) -> list[tuple]:
    return list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))


def _header_index(grid, first_label):
    return next(index for index, row in enumerate(grid) if row and row[0] == first_label)


def test_the_six_reports_are_registered():
    assert set(REPORTS) == {
        "absences",
        "balances",
        "waiting",
        "permission-hours",
        "remote-days",
        "certificates",
    }
    assert [name for name, definition in REPORTS.items() if definition.hr_only] == ["certificates"]


def test_unknown_report_is_not_found(world):
    with pytest.raises(NotFound) as refused:
        run_report(user=world.hr_user, name="payroll", query={})

    assert refused.value.code == "reports.unknown_report"


def test_run_returns_definition_parameters_and_result(world, november):
    run = run_report(user=world.hr_user, name="absences", query=NOVEMBER)

    assert run.definition.name == "absences"
    assert run.params.scope == "all"
    assert len(run.result.rows) == 7


def test_a_manager_runs_reports_for_the_team(world, november):
    run = run_report(user=world.manager_user, name="absences", query=NOVEMBER)

    assert run.params.scope == "team"
    assert len(run.result.rows) == 6


def test_certificates_are_for_hr_only(world, november):
    for user in (world.manager_user, world.exec_user):
        with pytest.raises(NotAllowed) as refused:
            run_report(user=user, name="certificates", query={})
        assert refused.value.code == "reports.hr_only"

    assert run_report(user=world.hr_user, name="certificates", query={}).definition.hr_only is True


def test_export_returns_a_workbook_and_a_file_name(world, november):
    with translation.override("fr"):
        content, filename = export_report(user=world.hr_user, name="absences", query=NOVEMBER)

    assert filename == "absences_2026-12-01.xlsx"
    grid = _grid(content)
    header = _header_index(grid, "Salarié")
    assert grid[header][:7] == ("Salarié", "Département", "Type", "Du", "Au", "Jours", "Heures")
    assert grid[header + 8][0] == "Total"  # seven data rows, then the overall total
    assert grid[header + 8][5] == 8


def test_export_follows_the_readers_language(world, november):
    with translation.override("en"):
        content, _ = export_report(user=world.hr_user, name="absences", query=NOVEMBER)

    assert _grid(content)[_header_index(_grid(content), "Employee")][3:5] == ("From", "To")


def test_export_is_audited_with_its_filters(world, november):
    export_report(user=world.hr_user, name="absences", query={**NOVEMBER, "type": "sick"})

    event = AuditEvent.objects.filter(action="report.exported").latest("at")
    assert event.category == "business"
    assert event.actor == world.hr_user
    assert (event.object_type, event.object_id) == ("report", "absences")
    assert event.changes == {
        "filters": {"scope": "all", "from": "2026-11-01", "to": "2026-11-30", "type": "sick"},
        "rows": 2,
    }


def test_looking_at_a_report_is_not_an_export(world, november):
    run_report(user=world.hr_user, name="absences", query=NOVEMBER)

    assert not AuditEvent.objects.filter(action="report.exported").exists()


def test_filter_options_for_hr(world, november):
    with translation.override("fr"):
        options = filter_options(user=world.hr_user, scope=None)

    assert options["scope"] == "all"
    assert {
        "id": world.dept_a.id,
        "name": organisation_selectors.department_name(world.owner, "fr"),
    } in options["departments"]
    assert {
        "public_id": str(world.owner.public_id),
        "name": organisation_selectors.full_name(world.owner),
    } in options["employees"]
    assert {"code": "annual", "name": world.annual.name_fr} in options["types"]
    assert [step["code"] for step in options["steps"]] == ["line_manager", "hr_role", "named_person"]
    assert options["steps"][1]["name"] == "RH"
    assert {
        "public_id": str(world.manager_user.public_id),
        "name": organisation_selectors.full_name(world.manager),
    } in options["actors"]
    assert str(world.sysadmin_user.public_id) not in {actor["public_id"] for actor in options["actors"]}


def test_filter_options_for_a_manager_are_the_team_without_departments(world, november):
    options = filter_options(user=world.manager_user, scope="team")

    assert options["scope"] == "team"
    assert options["departments"] == []
    assert options["actors"] == []
    assert {employee["public_id"] for employee in options["employees"]} == {
        str(world.owner.public_id),
        str(world.peer.public_id),
    }


def test_filter_options_refuse_a_scope_the_user_does_not_hold(world):
    with pytest.raises(NotAllowed):
        filter_options(user=world.manager_user, scope="all")
