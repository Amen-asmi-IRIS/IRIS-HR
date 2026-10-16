from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from apps.reports.tabular.framework import Column, Period, ReportResult, TotalLine
from apps.reports.tabular.render import to_json, to_workbook

RESULT = ReportResult(
    columns=[
        Column("employee", "Salarié", "text"),
        Column("period", "Dates", "period"),
        Column("days", "Jours", "number"),
        Column("since", "Depuis", "datetime"),
        Column("due", "Échéance", "date"),
    ],
    rows=[
        {
            "employee": "Amal Trabelsi",
            "period": Period(date(2026, 11, 2), date(2026, 11, 3)),
            "days": Decimal("1.5"),
            "since": datetime(2026, 11, 2, 8, 30, tzinfo=UTC),
            "due": date(2026, 11, 4),
        },
        {
            "employee": "Sami Ben Ali",
            "period": Period(date(2026, 11, 9), date(2026, 11, 9)),
            "days": None,
            "since": None,
            "due": None,
        },
    ],
    totals=[
        TotalLine("overall", "Total", {"days": Decimal("1.5")}),
        TotalLine("by_employee", "Amal Trabelsi", {"days": Decimal("1.5")}),
    ],
)


def _grid(content: bytes) -> list[tuple]:
    return list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))


def _header_index(grid: list[tuple], first_label: str) -> int:
    return next(index for index, row in enumerate(grid) if row and row[0] == first_label)


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


def test_json_cells_are_plain_values():
    body = to_json(RESULT)

    assert body["columns"][1] == {"key": "period", "label": "Dates", "kind": "period"}
    first = body["rows"][0]
    assert first["employee"] == "Amal Trabelsi"
    assert first["period"] == {"start": "2026-11-02", "end": "2026-11-03"}
    assert first["days"] == "1.50"
    assert first["since"] == "2026-11-02T08:30:00+00:00"
    assert first["due"] == "2026-11-04"
    assert body["rows"][1]["days"] is None
    assert body["rows"][1]["since"] is None


def test_json_totals_are_lines_or_null():
    assert to_json(RESULT)["totals"] == {
        "lines": [
            {"group": "overall", "label": "Total", "values": {"days": "1.50"}},
            {"group": "by_employee", "label": "Amal Trabelsi", "values": {"days": "1.50"}},
        ]
    }
    assert to_json(ReportResult(columns=[], rows=[]))["totals"] is None


def test_workbook_splits_a_period_into_two_date_columns():
    grid = _grid(to_workbook(RESULT, title="Absences prises", lang="fr"))
    header = _header_index(grid, "Salarié")

    assert grid[header] == ("Salarié", "Du", "Au", "Jours", "Depuis", "Échéance")
    first = grid[header + 1]
    assert first[0] == "Amal Trabelsi"
    assert _as_date(first[1]) == date(2026, 11, 2)
    assert _as_date(first[2]) == date(2026, 11, 3)
    assert first[3] == 1.5
    assert _as_date(first[5]) == date(2026, 11, 4)


def test_workbook_datetimes_are_tunis_wall_clock_without_zone():
    grid = _grid(to_workbook(RESULT, title="Absences prises", lang="fr"))
    header = _header_index(grid, "Salarié")

    assert grid[header + 1][4] == datetime(2026, 11, 2, 9, 30)


def test_workbook_keeps_empty_cells_empty():
    grid = _grid(to_workbook(RESULT, title="Absences prises", lang="fr"))
    header = _header_index(grid, "Salarié")

    assert grid[header + 2][3] is None
    assert grid[header + 2][4] is None


def test_workbook_ends_with_the_total_lines():
    grid = _grid(to_workbook(RESULT, title="Absences prises", lang="fr"))
    header = _header_index(grid, "Salarié")

    overall, by_employee = grid[header + 3], grid[header + 4]
    assert (overall[0], overall[3]) == ("Total", 1.5)
    assert (by_employee[0], by_employee[3]) == ("Total – Amal Trabelsi", 1.5)


def test_english_headers_for_an_english_reader():
    grid = _grid(to_workbook(RESULT, title="Absences taken", lang="en"))
    header = _header_index(grid, "Salarié")

    assert grid[header][1:3] == ("From", "To")


NOTED = ReportResult(
    columns=[Column("employee", "Salarié", "text"), Column("balance", "Solde", "number")],
    rows=[{"employee": "Amal Trabelsi", "balance": Decimal("12.5")}],
    totals=[TotalLine("overall", "Total", {"balance": Decimal("12.5")})],
    notes=("Les soldes comprennent les crédits de fin de mois à venir.",),
)


def test_json_carries_the_notes_and_an_empty_list_without():
    assert to_json(NOTED)["notes"] == ["Les soldes comprennent les crédits de fin de mois à venir."]
    assert to_json(RESULT)["notes"] == []


def test_workbook_keeps_the_notes_on_a_sheet_of_their_own():
    workbook = load_workbook(BytesIO(to_workbook(NOTED, title="Soldes", lang="fr")))

    assert workbook.sheetnames == ["Soldes", "Notes"]
    assert list(workbook["Soldes"].iter_rows(values_only=True)) == [
        ("Salarié", "Solde"),
        ("Amal Trabelsi", 12.5),
        ("Total", 12.5),
    ]
    assert list(workbook["Notes"].iter_rows(values_only=True)) == [
        ("Les soldes comprennent les crédits de fin de mois à venir.",)
    ]


def test_workbook_without_notes_has_one_sheet():
    assert load_workbook(BytesIO(to_workbook(RESULT, title="Absences", lang="fr"))).sheetnames == ["Absences"]
