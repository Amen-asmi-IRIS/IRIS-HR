from datetime import date
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.organisation import selectors as organisation_selectors
from apps.reports.tabular.render import XLSX_MEDIA_TYPE
from apps.reports.tests import scenario

pytestmark = pytest.mark.django_db

NOVEMBER = {"from": "2026-11-01", "to": "2026-11-30"}


@pytest.fixture(autouse=True)
def french_readers(world):
    """Labels are asserted in French; the language of each account decides the language of the answer."""
    for user in (world.hr_user, world.manager_user, world.exec_user):
        user.language = "fr"
        user.save(update_fields=["language"])


def _report(name="absences"):
    return reverse("reports:report", kwargs={"name": name})


def test_a_report_as_json(world, november, client_for):
    response = client_for(world.hr_user).get(_report(), NOVEMBER)

    assert response.status_code == 200
    body = response.json()
    assert [column["key"] for column in body["columns"]] == [
        "employee",
        "department",
        "type",
        "period",
        "days",
        "hours",
    ]
    assert len(body["rows"]) == 7
    first = body["rows"][0]
    assert set(first["period"]) == {"start", "end"}
    assert body["totals"]["lines"][0] == {
        "group": "overall",
        "label": "Total",
        "values": {"days": "8.00", "hours": "2.50"},
    }
    assert body["notes"] == []


def test_a_report_as_a_workbook(world, november, client_for):
    response = client_for(world.hr_user).get(_report(), {**NOVEMBER, "format": "xlsx"})

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_MEDIA_TYPE
    assert response["Content-Disposition"] == 'attachment; filename="absences_2026-12-01.xlsx"'
    assert response["X-Content-Type-Options"] == "nosniff"
    assert load_workbook(BytesIO(response.content)).active.max_row >= 8
    assert AuditEvent.objects.filter(action="report.exported", object_id="absences").count() == 1


def test_json_does_not_write_an_export_event(world, november, client_for):
    client_for(world.hr_user).get(_report(), NOVEMBER)

    assert not AuditEvent.objects.filter(action="report.exported").exists()


def test_labels_follow_the_language_of_the_account(world, november, client_for):
    world.hr_user.language = "en"
    world.hr_user.save(update_fields=["language"])

    body = client_for(world.hr_user).get(_report(), NOVEMBER).json()

    assert body["columns"][0]["label"] == "Employee"


def test_unknown_report_is_404(world, client_for):
    response = client_for(world.hr_user).get(_report("payroll"))

    assert response.status_code == 404
    assert response.json()["code"] == "reports.unknown_report"


def test_a_missing_period_is_a_field_problem(world, client_for):
    response = client_for(world.hr_user).get(_report())

    assert response.status_code == 400
    body = response.json()
    assert body["code"] == "reports.period_required"
    assert body["errors"][0]["field"] == "from"
    assert body["errors"][0]["detail"]


def test_a_manager_gets_the_team(world, november, client_for):
    body = client_for(world.manager_user).get(_report(), NOVEMBER).json()

    assert {row["employee"] for row in body["rows"]} == {organisation_selectors.full_name(world.owner)}


def test_a_manager_cannot_ask_for_everyone(world, november, client_for):
    response = client_for(world.manager_user).get(_report(), {**NOVEMBER, "scope": "all"})

    assert response.status_code == 403
    assert response.json()["code"] == "reports.scope_not_allowed"


def test_a_manager_cannot_name_someone_outside_the_team(world, november, client_for):
    response = client_for(world.manager_user).get(
        _report(), {**NOVEMBER, "employee": str(world.outsider.public_id)}
    )

    assert response.status_code == 404
    assert response.json()["code"] == "reports.employee_not_found"


def test_an_employee_without_a_team_is_refused(world, november, client_for):
    assert client_for(world.owner_user).get(_report(), NOVEMBER).status_code == 403


def test_certificates_are_for_hr(world, november, client_for):
    assert client_for(world.hr_user).get(_report("certificates")).status_code == 200
    for user in (world.manager_user, world.exec_user):
        response = client_for(user).get(_report("certificates"))
        assert response.status_code == 403
        assert response.json()["code"] == "reports.hr_only"


def test_no_report_shows_the_sick_leave_comment(world, november, client_for):
    scenario.pending(
        world,
        employee=world.peer,
        request_type=world.sick,
        start=date(2026, 11, 26),
        comment=scenario.SICK_COMMENT,
    )
    client = client_for(world.hr_user)

    for name, query in (
        ("absences", NOVEMBER),
        ("balances", {}),
        ("waiting", {}),
        ("permission-hours", {}),
        ("remote-days", {}),
        ("certificates", {}),
    ):
        response = client.get(_report(name), query)
        assert response.status_code == 200, name
        assert "Karray" not in response.content.decode(), name
        workbook = client.get(_report(name), {**query, "format": "xlsx"})
        cells = load_workbook(BytesIO(workbook.content)).active.iter_rows(values_only=True)
        assert "Karray" not in " ".join(str(cell) for row in cells for cell in row if cell is not None), name


def test_filter_options(world, november, client_for):
    url = reverse("reports:report-filters")

    hr = client_for(world.hr_user).get(url).json()
    manager = client_for(world.manager_user).get(url, {"scope": "team"}).json()

    assert hr["scope"] == "all"
    assert {"code": "sick", "name": world.sick.name_fr} in hr["types"]
    assert str(world.manager_user.public_id) in {actor["public_id"] for actor in hr["actors"]}
    assert manager["scope"] == "team"
    assert manager["departments"] == []
    assert manager["actors"] == []
    assert len(manager["employees"]) == 2


def test_kpmg_export_downloads_a_workbook(world, november, client_for):
    response = client_for(world.hr_user).get(reverse("reports:kpmg-export"), NOVEMBER)

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_MEDIA_TYPE
    assert response["Content-Disposition"] == 'attachment; filename="kpmg_2026-11-01_2026-11-30.xlsx"'
    assert response["Cache-Control"] == "no-store"
    assert AuditEvent.objects.filter(action="export.kpmg").count() == 1


def test_kpmg_export_with_nothing_to_export(world, november, client_for):
    response = client_for(world.hr_user).get(
        reverse("reports:kpmg-export"), {**NOVEMBER, "department": "999999"}
    )

    assert response.status_code == 400
    assert response.json()["code"] == "reports.nothing_to_export"
    assert not AuditEvent.objects.filter(action="export.kpmg").exists()


def test_kpmg_export_is_for_hr_only(world, november, client_for, anon_client):
    url = reverse("reports:kpmg-export")

    assert anon_client.get(url, NOVEMBER).status_code == 401
    for user in (world.exec_user, world.manager_user, world.owner_user, world.sysadmin_user):
        assert client_for(user).get(url, NOVEMBER).status_code == 403


def test_kpmg_default_period(world, november, client_for):
    response = client_for(world.hr_user).get(reverse("reports:kpmg-default-period"))

    assert response.status_code == 200
    assert response.json() == {"start": "2026-11-01", "end": "2026-11-30", "cutoff_day": None}
